# -*- coding: utf-8 -*-
"""Security-sensitive ONLYOFFICE file, token and editor helpers."""
import hashlib
import os
import shutil
import tempfile
import uuid
import zipfile
from datetime import datetime
from urllib.parse import urlparse

try:
    import jwt
except ImportError:  # application can still start and legacy documents remain usable
    jwt = None
try:
    import requests
except ImportError:
    requests = None
from flask import current_app, url_for
from itsdangerous import BadSignature, SignatureExpired, URLSafeTimedSerializer

from app.extensions import db
from app.models import DocVersion

OFFICE_TYPES = {'docx': 'word', 'xlsx': 'cell', 'pptx': 'slide'}
MIME_TYPES = {
    'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    'pptx': 'application/vnd.openxmlformats-officedocument.presentationml.presentation'
}
REQUIRED_PARTS = {'docx': 'word/document.xml', 'xlsx': 'xl/workbook.xml',
                  'pptx': 'ppt/presentation.xml'}


def normalize_server_url(value=None):
    return (value if value is not None else current_app.config.get('ONLYOFFICE_SERVER_URL', '')).strip().rstrip('/')


def is_onlyoffice_configured():
    cfg = current_app.config
    return bool(cfg.get('ONLYOFFICE_ENABLED') and normalize_server_url() and jwt is not None and requests is not None and
                (not cfg.get('ONLYOFFICE_JWT_ENABLED') or cfg.get('ONLYOFFICE_JWT_SECRET')))


def office_type_from_extension(ext):
    return OFFICE_TYPES.get((ext or '').lower().lstrip('.'))


def checksum_file(path):
    digest = hashlib.sha256()
    with open(path, 'rb') as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b''):
            digest.update(chunk)
    return digest.hexdigest()


def safe_document_path(relative_path, version=False):
    root = current_app.config['DOCUMENT_VERSION_FOLDER' if version else 'DOCUMENT_STORAGE_FOLDER']
    root = os.path.realpath(root)
    relative_path = (relative_path or '').replace('\\', '/')
    if not relative_path or os.path.isabs(relative_path) or '..' in relative_path.split('/'):
        raise ValueError('Invalid document path')
    result = os.path.realpath(os.path.join(root, *relative_path.split('/')))
    if os.path.commonpath([root, result]) != root:
        raise ValueError('Document path escapes storage')
    return result


def validate_office_file(path, ext, max_bytes=None):
    ext = (ext or '').lower().lstrip('.')
    if ext not in REQUIRED_PARTS:
        raise ValueError('仅支持 docx、xlsx、pptx 文件')
    size = os.path.getsize(path)
    limit = max_bytes or current_app.config['DOCUMENT_MAX_UPLOAD_MB'] * 1024 * 1024
    if size <= 0 or size > limit:
        raise ValueError('文件为空或超过大小限制')
    if not zipfile.is_zipfile(path):
        raise ValueError('文件内容不是有效的 Office Open XML 格式')
    total_uncompressed = 0
    with zipfile.ZipFile(path) as archive:
        names = set()
        for info in archive.infolist():
            name = info.filename.replace('\\', '/')
            if name.startswith('/') or '..' in name.split('/'):
                raise ValueError('Office 文件包含不安全路径')
            total_uncompressed += info.file_size
            if total_uncompressed > max(limit * 20, 200 * 1024 * 1024):
                raise ValueError('Office 文件解压后体积异常')
            names.add(name)
        if REQUIRED_PARTS[ext] not in names:
            raise ValueError('文件扩展名与 Office 内容不匹配')
    return size, checksum_file(path)


def create_blank_office_file(path, ext):
    ext = ext.lower().lstrip('.')
    if ext == 'docx':
        from docx import Document
        Document().save(path)
    elif ext == 'xlsx':
        from openpyxl import Workbook
        book = Workbook()
        book.save(path)
    elif ext == 'pptx':
        from pptx import Presentation
        Presentation().save(path)
    else:
        raise ValueError('不支持的文档类型')
    return validate_office_file(path, ext)


def build_document_key(doc_id, file_version, checksum):
    raw = '{}:{}:{}'.format(int(doc_id), int(file_version or 1), checksum or '')
    return 'doc-{}-v{}-{}'.format(int(doc_id), int(file_version or 1), hashlib.sha256(raw.encode()).hexdigest()[:24])


def _serializer():
    return URLSafeTimedSerializer(current_app.config['SECRET_KEY'], salt='onlyoffice-content-v1')


def build_content_token(doc_id, purpose='content'):
    return _serializer().dumps({'doc_id': int(doc_id), 'purpose': purpose})


def verify_content_token(token, doc_id, purpose='content'):
    try:
        data = _serializer().loads(token, max_age=current_app.config['DOCUMENT_URL_TOKEN_MAX_AGE'])
        return data.get('doc_id') == int(doc_id) and data.get('purpose') == purpose
    except (BadSignature, SignatureExpired, TypeError, ValueError):
        return False


def sign_editor_config(config):
    if not current_app.config.get('ONLYOFFICE_JWT_ENABLED'):
        return ''
    if jwt is None: raise RuntimeError('PyJWT is not installed')
    return jwt.encode(config, current_app.config['ONLYOFFICE_JWT_SECRET'], algorithm='HS256')


def verify_callback_token(token):
    if not current_app.config.get('ONLYOFFICE_JWT_ENABLED'):
        return None
    if jwt is None: raise ValueError('PyJWT is not installed')
    if not token:
        raise ValueError('Missing callback JWT')
    try:
        return jwt.decode(token, current_app.config['ONLYOFFICE_JWT_SECRET'], algorithms=['HS256'])
    except jwt.PyJWTError as exc:
        raise ValueError('Invalid callback JWT') from exc


def build_editor_config(doc, user, can_edit):
    mode = 'edit' if can_edit else 'view'
    public = current_app.config.get('APP_PUBLIC_URL') or ''
    content_url = public + url_for('docs.document_content', doc_id=doc.id,
                                   token=build_content_token(doc.id, 'content'))
    callback_url = public + url_for('docs.onlyoffice_callback', doc_id=doc.id,
                                    callback_token=build_content_token(doc.id, 'callback'))
    config = {
        'documentType': doc.office_type,
        'document': {'fileType': doc.file_ext, 'key': doc.document_key,
                     'title': doc.original_filename or doc.title,
                     'url': content_url,
                     'permissions': {'edit': bool(can_edit), 'download': True, 'print': True}},
        'editorConfig': {'mode': mode, 'lang': 'zh-CN', 'callbackUrl': callback_url,
                         'user': {'id': str(user.id), 'name': user.real_name or user.username},
                         'customization': {'autosave': True, 'forcesave': True}}
    }
    token = sign_editor_config(config)
    if token:
        config['token'] = token
    return config


def create_file_version(doc, user_id, summary):
    source = safe_document_path(doc.storage_relpath)
    if not os.path.isfile(source):
        raise IOError('Current document file is missing')
    max_ver = db.session.query(db.func.max(DocVersion.version_num)).filter_by(doc_id=doc.id).scalar() or 0
    relpath = '{}/{}-v{}-{}.{}'.format(doc.id, doc.id, max_ver + 1, uuid.uuid4().hex, doc.file_ext)
    target = safe_document_path(relpath, version=True)
    os.makedirs(os.path.dirname(target), exist_ok=True)
    shutil.copy2(source, target)
    version = DocVersion(doc_id=doc.id, content='', title=doc.title, version_num=max_ver + 1,
                         change_summary=summary, created_by=user_id, storage_relpath=relpath,
                         file_size=os.path.getsize(target), checksum=checksum_file(target),
                         version_kind='onlyoffice')
    db.session.add(version)
    return version


def restore_file_version(doc, version, user_id):
    historical = safe_document_path(version.storage_relpath, version=True)
    if not os.path.isfile(historical):
        raise IOError('历史版本文件不存在')
    validate_office_file(historical, doc.file_ext)
    current = safe_document_path(doc.storage_relpath)
    create_file_version(doc, user_id, '恢复前自动保存')
    fd, temporary = tempfile.mkstemp(dir=os.path.dirname(current), suffix='.tmp')
    os.close(fd)
    try:
        shutil.copy2(historical, temporary)
        os.replace(temporary, current)
    finally:
        if os.path.exists(temporary): os.remove(temporary)
    doc.file_size = os.path.getsize(current)
    doc.checksum = checksum_file(current)
    doc.file_version = (doc.file_version or 1) + 1
    doc.document_key = build_document_key(doc.id, doc.file_version, doc.checksum)
    doc.last_editor_id = user_id
    doc.last_saved_at = datetime.utcnow()
    doc.updated_at = datetime.utcnow()


def save_callback_file(doc, download_url, user_id, final_save=True):
    parsed = urlparse(download_url or '')
    allowed = current_app.config.get('ONLYOFFICE_DOWNLOAD_HOSTS', ())
    if parsed.scheme not in ('http', 'https') or not parsed.hostname or parsed.hostname.lower() not in allowed:
        raise ValueError('Callback download host is not allowed')
    target = safe_document_path(doc.storage_relpath)
    limit = current_app.config['DOCUMENT_MAX_UPLOAD_MB'] * 1024 * 1024
    if requests is None: raise RuntimeError('requests is not installed')
    response = requests.get(download_url, stream=True, timeout=(5, 30), allow_redirects=False)
    response.raise_for_status()
    fd, temporary = tempfile.mkstemp(dir=os.path.dirname(target), suffix='.' + doc.file_ext)
    os.close(fd)
    try:
        size = 0
        with open(temporary, 'wb') as output:
            for chunk in response.iter_content(1024 * 1024):
                if not chunk: continue
                size += len(chunk)
                if size > limit: raise ValueError('Callback file exceeds size limit')
                output.write(chunk)
        size, checksum = validate_office_file(temporary, doc.file_ext, limit)
        if checksum == doc.checksum:  # idempotent repeated callback
            return False
        create_file_version(doc, user_id, 'ONLYOFFICE 保存前快照')
        os.replace(temporary, target)
        doc.file_size, doc.checksum = size, checksum
        if final_save:
            doc.file_version = (doc.file_version or 1) + 1
            doc.document_key = build_document_key(doc.id, doc.file_version, checksum)
        doc.last_editor_id = user_id
        doc.last_saved_at = datetime.utcnow()
        doc.updated_at = datetime.utcnow()
        return True
    finally:
        response.close()
        if os.path.exists(temporary): os.remove(temporary)
